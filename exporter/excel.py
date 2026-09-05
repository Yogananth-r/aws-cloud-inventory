from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from config import DEFAULT_REPORT


HEADER_FILL = PatternFill(
    fill_type="solid",
    start_color="1F4E78",
    end_color="1F4E78",
)


class ExcelExporter:
    """Export inventory data into an Excel workbook."""

    def __init__(self, output_file=DEFAULT_REPORT):
        self.output_file = output_file
        self.workbook = Workbook()

    def write_sheet(self, sheet_name: str, data: list[dict]):
        """Create a worksheet from a list of dictionaries."""
        sheet = self.workbook.active if len(self.workbook.sheetnames) == 1 else self.workbook.create_sheet()
        sheet.title = sheet_name

        if not data:
            sheet.append(["No resources found"])
            return

        headers = list(data[0].keys())
        sheet.append(headers)

        # Header formatting
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL

        # Data rows
        for row in data:
            sheet.append(list(row.values()))

        # Auto-size columns
        for column in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column)
            column_letter = get_column_letter(column[0].column)
            sheet.column_dimensions[column_letter].width = min(max_length + 2, 40)

        # Freeze header row
        sheet.freeze_panes = "A2"

        # Enable filters
        sheet.auto_filter.ref = sheet.dimensions

    def save(self):
        self.workbook.save(self.output_file)
        return self.output_file